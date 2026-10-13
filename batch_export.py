"""Write large projects to Excel a record at a time, with the pilot's markers."""
import re
import uuid
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

COLORS={'uncertain':'FFF0CC','unmapped':'FDE2DE','confirmed':'E2F2E8','generated':'EDF2F8'}


def write_workbook(path, state, export_record, headers):
    workbook=Workbook(write_only=True)
    font=Font(name='Arial',size=11)
    header_font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
    alignment=Alignment(vertical='top',wrap_text=True)
    fills={k:PatternFill('solid',fgColor=v) for k,v in COLORS.items()}

    def table(name, columns):
        sheet=workbook.create_sheet(name)
        sheet.freeze_panes='C2' if name=='Converted records' else 'B2'
        sheet.sheet_view.showGridLines=False
        sheet.sheet_format.defaultRowHeight=60
        for i in range(len(columns)):sheet.column_dimensions[get_column_letter(i+1)].width=25
        sheet.row_dimensions[1].height=42
        cells=[]
        for value in columns:
            cell=WriteOnlyCell(sheet,value=value)
            cell.font=header_font;cell.fill=PatternFill('solid',fgColor='243B53');cell.alignment=alignment
            cells.append(cell)
        sheet.append(cells)
        return sheet

    if state.get('partial_export'):
        summary=table('Export summary',['PARTIAL EXPORT','Scope'])
        summary.column_dimensions['B'].width=100
        def note(label,value):
            cells=[]
            for text in (label,value):
                cell=WriteOnlyCell(summary,value=text);cell.data_type='s';cell.font=font;cell.alignment=alignment
                cells.append(cell)
            summary.append(cells)
        note('Available results only',state['export_scope'])
        note('Review still required','Unfinished records may be incomplete. [i/l] marks alternatives and [?] marks unreadable or unlocated text.')
        note('Source images','The complete upload order is listed below.')
        done=state.get('completed_pages',0)
        included={f.get('file') for record in state['records'] for f in record.get('fragments',[])}
        for i,page in enumerate(state['pages']):
            note(page['name'],'Fully scanned' if i<done else 'Partly scanned' if page['file'] in included else 'Not yet scanned')
    converted=table('Converted records',headers)
    converted.column_dimensions['A'].width=42
    for c in (14,26,37):converted.column_dimensions[get_column_letter(c)].width=55
    raw=table('Text regions',['Record key','Source image','Region','Transcription with markers'])
    review=table('Review status',['Record key','Field or review ID','Category','Original reading or issue','Confirmed text / explanation','Status'])
    fragments=table('Source fragments',['Record key','Source image','Fragment','Assignment','Source coordinates','Layout notes'])

    def append(sheet,values,statuses=None):
        cells=[]
        for i,value in enumerate(values):
            cell=WriteOnlyCell(sheet,value=value)
            # Literal OCR/user text must never become an Excel formula.
            if isinstance(value,str):cell.data_type='s'
            cell.number_format='@';cell.font=font;cell.alignment=alignment
            if statuses:
                if statuses[i] in fills:cell.fill=fills[statuses[i]]
                if isinstance(value,str) and '[' not in value:
                    number=None
                    if i in (18,21,23,24) and re.fullmatch(r'\$\d[\d,]*\.\d{2}',value):
                        number=float(value.replace('$','').replace(',',''));cell.number_format='"$"#,##0.00'
                    elif i==19 and re.fullmatch(r'\d+ Months',value):
                        number=int(value[:-7]);cell.number_format='0" Months"'
                    elif i==20 and re.fullmatch(r'\d+(\.\d+)?%',value):
                        number=float(value[:-1])/100;cell.number_format='0.########%'
                    elif i in (2,9) and re.fullmatch(r'\d{2}/\d{2}/\d{4}',value):
                        try:cell.value=datetime.strptime(value,'%m/%d/%Y');cell.number_format='mm/dd/yyyy'
                        except ValueError:pass
                    if number is not None and abs(number)<1e14:cell.value=number;cell.data_type='n'
            cells.append(cell)
        sheet.append(cells)

    for record in state['records']:
        view=dict(state,records=[record],page_issues=[])
        for key in ('field_regions','field_decisions'):
            view[key]={k:v for k,v in state.get(key,{}).items() if k.rsplit(':',1)[0]==record['key']}
        payload=export_record(view)
        for i,row in enumerate(payload['rows']):append(converted,row,payload.get('statuses',[None]*len(payload['rows']))[i])
        for row in payload['lines']:append(raw,row)
        for row in payload['review']:append(review,row)
        for row in payload.get('fragments',[]):append(fragments,row)
    for issue in state.get('page_issues',[]):append(review,['','','Page review',issue,'','Needs review'])
    append(review,[])
    for row in [
        ['Legend','Amber: character ambiguity or field rule needs review.'],
        ['Legend','Red: field not located. Green: user-confirmed; any rule warning is retained.'],
        ['Legend','Blue-grey: generated metadata or remarks. White: unreviewed OCR.'],
        ['Markers','[i/l] = alternatives; [?] = unreadable or unlocated text.'],
        ['Source',state.get('export_scope','All project records are exported, including records outside the current screen.')]
    ]:append(review,row)
    temporary=path.with_name('export-'+uuid.uuid4().hex+'.xlsx')
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:temporary.unlink(missing_ok=True)
