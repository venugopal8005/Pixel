"""Concurrent review, continuation boundaries, cancellation and checkpoint recovery."""
import copy
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook
import review_app as app
from scan_state import prepare_resume, line_counter, export_snapshot
from test_batch_upload import QuietHandler
from field_crop import target


def block(n=0, opened=False, first=True):
    return dict(box=[0, n*40, 100, (n+1)*40], column=0, block=n,
                first_in_column=first, open_bottom=opened)


class ProgressiveScanTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='pixel-progressive-')
        self.previous=app.DATA;app.DATA=Path(self.tmp.name)
        self.state=dict(id='d'*32, project_name='Progressive test', schema_version=2,
            status='scanning', progress='Waiting', pages=[dict(name=f'AQ{i+1}.png',file=f'page-{i}.png') for i in range(2)],
            records=[],decisions={},history=[],progressive_scan=True,completed_pages=0)
        app.folder(self.state['id']).mkdir()
        for page in self.state['pages']:Image.new('RGB',(100,100),'white').save(app.folder(self.state['id'])/page['file'])
        app.save(self.state)
        self.server=app.ThreadingHTTPServer(('127.0.0.1',0),QuietHandler)
        self.http=threading.Thread(target=self.server.serve_forever,daemon=True);self.http.start()
        self.pause=threading.Event();self.release=threading.Event();self.worker=None

    def tearDown(self):
        self.release.set()
        if self.worker:self.worker.join(10)
        self.server.shutdown();self.server.server_close();self.http.join()
        app.LIVE_PROGRESS.pop(self.state['id'],None);app.SCAN_CANCEL.discard(self.state['id'])
        app.EXPORT_LOCKS.pop(self.state['id'],None)
        app.DATA=self.previous;self.tmp.cleanup()

    def request(self,path,data=None):
        req=urllib.request.Request(f'http://127.0.0.1:{self.server.server_port}'+path,
            data=None if data is None else json.dumps(dict(job=self.state['id'],**data)).encode(),
            headers={'Content-Type':'application/json','X-Review-Token':app.TOKEN})
        with urllib.request.urlopen(req,timeout=15) as response:return json.load(response)

    def start(self):
        self.worker=threading.Thread(target=app.scan,args=(self.state,),daemon=True);self.worker.start()
        self.assertTrue(self.pause.wait(8),'Scanner did not reach the controlled pause')

    def finish(self):
        self.release.set();self.worker.join(8);self.assertFalse(self.worker.is_alive())
        return app.load(self.state['id'])

    @staticmethod
    def engine(*args,**kwargs):return ([[[5,5],[85,5],[85,20],[5,20]]],0)

    @staticmethod
    def text(*args):return [dict(text='A',options=['A','B'],category='ambiguous',reason='test',x=0)]

    def test_page_one_review_and_export_during_page_two_preserve_every_review_kind(self):
        calls=[]
        def layout(image):
            calls.append(image)
            if len(calls)==2:self.pause.set();self.release.wait(8)
            return [block(),block(1,opened=True,first=False)] if len(calls)==1 else [block()]
        with patch('rapidocr_onnxruntime.RapidOCR',return_value=self.engine),patch('detect_records.blocks',side_effect=layout),\
             patch.object(app,'read_header',side_effect=[('1234567890','1234567890'),('1234567891','1234567891'),(None,'continuation')]),\
             patch.object(app,'read_line',side_effect=self.text):
            self.start()
            view=self.request('/api/workspace?job='+self.state['id'])
            self.assertEqual(view['state']['completed_pages'],1)
            self.assertEqual(len(view['rows']),1)
            self.assertEqual(view['state']['pending_record_total'],1)
            self.assertEqual(view['state']['image_navigation'][0]['status'],'continuation')
            before=copy.deepcopy(app.load(self.state['id'])['records'][0])
            self.request('/api/confirm',dict(segment='l1-s0',value='B'))
            self.request('/api/field',dict(record='record-1',field=6,value='Verified address'))
            # Crop OCR stores its lines inside the field selection; a page checkpoint must keep them too.
            with app.LOCK:
                edited=app.load(self.state['id'])
                target(edited,'record-1',4)
                edited['field_regions']={'record-1:4':dict(kind='rectangles',id='saved-crop',regions=[],lines=[],first='',last='')}
                edited['field_crop_history']=[dict(field='record-1:4',selection={'id':'prior-crop'})]
                edited['verified_lines']={'l1':True};app.save(edited)
            with self.assertRaises(urllib.error.HTTPError):
                self.request('/api/field',dict(record='record-2',field=6,value='Too early'))
            with self.assertRaises(urllib.error.HTTPError):
                self.request('/api/boundary',dict(record='record-1',action='join_previous'))
            # An incomplete field selection is irrelevant to this test's export; keep a valid line selection instead.
            with app.LOCK:
                edited=app.load(self.state['id']);edited['field_regions']['record-1:4']={'first':'l1-s0','last':'l1-s0'};app.save(edited)
            result=self.request('/api/export',{})
            self.assertTrue(result['partial']);self.assertEqual(result['records'],1)
            self.assertIn('partial=1',result['url'])
            book=load_workbook(app.folder(self.state['id'])/'converted.xlsx',read_only=True)
            self.assertEqual(book.sheetnames[0],'Export summary')
            self.assertEqual(len(list(book['Converted records'].values)),2)
            self.assertTrue(any('PARTIAL EXPORT' in str(row) for row in book['Review status'].values));book.close()
            final=self.finish()
        self.assertEqual(final['status'],'ready');self.assertEqual(final['completed_pages'],2)
        self.assertEqual(len(final['records']),2,'Continuation must not become a duplicate record')
        self.assertEqual(len(final['records'][1]['fragments']),2)
        self.assertEqual(final['records'][0],before,'Published OCR must stay stable while reviewed')
        self.assertEqual(final['decisions']['l1-s0'],'B')
        self.assertEqual(final['field_decisions']['record-1:6'],'Verified address')
        self.assertIn('record-1:4',final['field_regions']);self.assertEqual(final['field_crop_history'][0]['selection']['id'],'prior-crop')
        self.assertEqual(final['verified_lines'],{'l1':True});self.assertGreaterEqual(len(final['history']),2)

    def test_cancel_queued_does_not_initialize_ocr_or_delete_images(self):
        app.SCAN_LOCK.acquire()
        try:
            with patch('rapidocr_onnxruntime.RapidOCR') as engine:
                self.worker=threading.Thread(target=app.scan,args=(self.state,),daemon=True);self.worker.start()
                result=self.request('/api/cancel',{})
                self.assertEqual(result['scan_stats']['stage'],'cancelling')
                self.worker.join(3);self.assertFalse(self.worker.is_alive());engine.assert_not_called()
        finally:app.SCAN_LOCK.release()
        self.assertEqual(app.load(self.state['id'])['status'],'cancelled')
        self.assertTrue((app.folder(self.state['id'])/'page-1.png').exists())
        self.assertEqual(self.request('/api/cancel',{})['status'],'cancelled')

    def test_cancel_running_keeps_current_reading_and_releases_scanner(self):
        def read(*args):self.pause.set();self.release.wait(8);return self.text()
        with patch('rapidocr_onnxruntime.RapidOCR',return_value=self.engine),patch('detect_records.blocks',return_value=[block()]),\
             patch.object(app,'read_header',return_value=('1234567890','1234567890')),patch.object(app,'read_line',side_effect=read) as reader:
            self.start();self.request('/api/cancel',{});final=self.finish()
            self.assertEqual(reader.call_count,1)
        self.assertEqual(final['status'],'cancelled');self.assertEqual(final['completed_pages'],0)
        self.assertEqual(len(final['records'][0]['lines']),1)
        self.assertEqual(len(app.workspace_page(final)['rows']),1)
        self.assertTrue(final['records'][0]['scan_partial'])
        self.assertTrue(app.SCAN_LOCK.acquire(timeout=1));app.SCAN_LOCK.release()

    def test_resume_uses_completed_page_cursor_and_keeps_ids_and_corrections(self):
        record=dict(key='record-1',record_id='1234567890',notes=[],review_ready=False,lines=[dict(id='l4',segments=[],fragment='p0-b0')],
            fragments=[dict(id='p0-b0',page='AQ1.png',file='page-0.png',box=[0,0,100,40],column=0,open_bottom=True)])
        self.state.update(records=[record],completed_pages=1,decisions={'l4-s0':'B'})
        app.save(self.state)
        with patch('rapidocr_onnxruntime.RapidOCR',return_value=self.engine),patch('detect_records.blocks',return_value=[block()]) as layout,\
             patch.object(app,'read_header',return_value=(None,'continued')),patch.object(app,'read_line',side_effect=self.text):
            app.scan(self.state)
        final=app.load(self.state['id']);self.assertEqual(final['status'],'ready')
        self.assertEqual(layout.call_count,1);self.assertEqual(len(final['records']),1)
        self.assertEqual([l['id'] for l in final['records'][0]['lines']],['l4','l5'])
        self.assertEqual(final['decisions'],{'l4-s0':'B'})


class LegacyRecoveryTests(unittest.TestCase):
    def fixture(self):
        return dict(status='scanning',pages=[dict(file=f'page-{i}.png') for i in range(3)],
            records=[dict(key=f'record-{i+1}',record_id=str(i),lines=[dict(id=f'l{i+1}',fragment=f'p{i}')],
                          fragments=[dict(file=f'page-{i}.png',id=f'p{i}',open_bottom=i==1)]) for i in range(3)],
            scan_stats=dict(page=3,block=2,page_blocks=4))

    def test_legacy_partial_page_replayed_but_finished_pages_kept(self):
        state=self.fixture();self.assertEqual(prepare_resume(state),2)
        self.assertEqual(len(state['records']),2);self.assertTrue(state['records'][0]['review_ready'])
        self.assertFalse(state['records'][1]['review_ready']);self.assertEqual(line_counter(state),2)

    def test_legacy_finished_last_page_is_not_replayed(self):
        state=self.fixture();state['scan_stats']['block']=4
        self.assertEqual(prepare_resume(state),3);self.assertEqual(len(state['records']),3)

    def test_export_without_ready_records_is_explicit_error(self):
        state=self.fixture();state.update(records=[],completed_pages=0,progressive_scan=True)
        with self.assertRaisesRegex(ValueError,'No records'):export_snapshot(state)


if __name__=='__main__':unittest.main()
