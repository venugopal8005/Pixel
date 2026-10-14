"""400-page transport/queue tests use synthetic images, not OCR accuracy claims."""
import base64
import copy
import io
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook
import review_app as app
from batch_upload import BatchUpload
from test_sheet_review import fixture


class QuietHandler(app.Handler):
    def log_message(self, *args):
        pass


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='pixel-batch-')
        self.previous = app.DATA
        app.DATA = Path(self.tmp.name)
        self.queued = []
        self.service = BatchUpload(app.load, app.save, app.folder, app.LOCK, self.queued.append)
        self.upload_patch = patch.object(app, 'UPLOADS', self.service)
        self.upload_patch.start()
        self.server = app.ThreadingHTTPServer(('127.0.0.1', 0), QuietHandler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        self.upload_patch.stop()
        app.DATA = self.previous
        self.tmp.cleanup()

    def request(self, path, data=None):
        req = urllib.request.Request(f'http://127.0.0.1:{self.server.server_port}'+path,
            data=None if data is None else json.dumps(data).encode(),
            headers={'Content-Type':'application/json','X-Review-Token':app.TOKEN})
        with urllib.request.urlopen(req, timeout=30) as response:
            return json.load(response)

    def images(self, count):
        result = []
        for i in range(count):
            raw = io.BytesIO()
            Image.new('RGB', (20,20), (i%256,i//256,30)).save(raw, format='PNG')
            result.append(dict(name=f'AQ{i+1}.png', size=len(raw.getvalue()),
                               data=base64.b64encode(raw.getvalue()).decode()))
        return result

    def create(self, images):
        return self.request('/api/upload/create', dict(name='Batch test',
            files=[dict(name=p['name'],size=p['size']) for p in images]))['id']

    def test_400_images_saved_in_order_and_queued_only_once_when_complete(self):
        images = self.images(400)
        job = self.create(images)
        with self.assertRaises(urllib.error.HTTPError):
            self.request('/api/upload/finish', dict(job=job))
        for i, page in enumerate(images):
            response = self.request('/api/upload/page', dict(job=job,index=i,data=page['data']))
            self.assertEqual(response['uploaded'], i+1)
        self.assertFalse(self.queued)
        self.assertEqual(len(list(app.folder(job).glob('page-*.png'))), 400)
        self.request('/api/upload/finish', dict(job=job))
        self.request('/api/upload/finish', dict(job=job))
        self.assertEqual(len(self.queued), 1)
        state = app.load(job)
        self.assertEqual([p['name'] for p in state['pages']], [p['name'] for p in images])
        self.assertEqual(state['status'], 'scanning')
        self.assertEqual(self.request('/api/jobs')[0]['pages'], 400)
        # Exercise the real page loop/progress/checkpoints with fast deterministic OCR.
        block=dict(box=[0,0,20,20],column=0,block=0,first_in_column=True,open_bottom=False)
        with patch('rapidocr_onnxruntime.RapidOCR',return_value=lambda *a,**kw:([],0)), \
             patch('detect_records.blocks',return_value=[block]), \
             patch.object(app,'read_header',side_effect=[(str(1800000000+i),'header') for i in range(400)]):
            app.scan(state)
        final=app.load(job)
        self.assertEqual(final['status'],'ready')
        self.assertEqual(len(final['records']),400)
        self.assertEqual(app.get_progress(job)['scan_stats']['page'],400)
        result=self.request('/api/export',dict(job=job))
        with urllib.request.urlopen(f'http://127.0.0.1:{self.server.server_port}'+result['url']) as response:
            book=load_workbook(io.BytesIO(response.read()),read_only=True)
        rows=list(book['Converted records'].values)
        self.assertEqual(len(rows),401)
        self.assertEqual(rows[1][0],'AQ1.png')
        self.assertEqual(rows[-1][0],'AQ400.png')
        self.assertEqual(len(list(book['Source fragments'].values)),401)
        book.close()
        app.LIVE_PROGRESS.pop(job,None)

    def test_restart_retry_duplicate_and_order_checks(self):
        images = self.images(3)
        job = self.create(images)
        first = dict(job=job,index=0,data=images[0]['data'])
        self.service.page(first)
        self.service = BatchUpload(app.load, app.save, app.folder, app.LOCK, self.queued.append)
        self.assertEqual(self.service.page(first)['uploaded'], 1)
        with self.assertRaises(ValueError):
            self.service.page(dict(job=job,index=0,data=images[1]['data']))
        with self.assertRaises(ValueError):
            self.service.page(dict(job=job,index=2,data=images[2]['data']))
        with self.assertRaises(ValueError):
            self.service.page(dict(job=job,index=1,data=images[0]['data']))
        self.assertEqual(len(app.load(job)['pages']),1)
        for i in (1,2):self.service.page(dict(job=job,index=i,data=images[i]['data']))
        self.service.finish(job)
        self.assertEqual(len(self.queued[0]['pages']),3)

    def test_limits_and_invalid_images_do_not_create_scans(self):
        for count in (0,401):
            with self.assertRaises(urllib.error.HTTPError):self.create(self.images(count))
        for size in (0,25_000_001):
            with self.assertRaises(urllib.error.HTTPError):
                self.create([dict(name='large.png',size=size)])
        images = self.images(1)
        with self.assertRaises(urllib.error.HTTPError):self.create(images*2)
        self.assertEqual(self.request('/api/jobs'),[])
        job = self.create(images)
        for data in ('not-base64',base64.b64encode(b'not an image').decode()):
            with self.assertRaises(urllib.error.HTTPError):
                self.request('/api/upload/page',dict(job=job,index=0,data=data))
        self.assertFalse(app.load(job)['pages'])
        self.assertFalse(self.queued)

    def test_workspace_only_maps_50_records_and_keeps_last_record_edit(self):
        state=fixture()
        template=state['records'][0]
        state['records']=[]
        for i in range(103):
            record=copy.deepcopy(template)
            record['key']=f'record-{i}'
            for line in record['lines']:
                line['id']=f'{i}-'+line['id']
                for segment in line['segments']:segment['id']=f'{i}-'+segment['id']
            state['records'].append(record)
        state['field_decisions']={'record-102:6':'Last record correction','record-0:6':'First record correction'}
        app.folder(state['id']).mkdir()
        app.save(state)
        result=self.request('/api/workspace?job='+state['id']+'&offset=100')
        self.assertEqual(len(result['state']['records']),3)
        self.assertEqual(len(result['rows']),3)
        self.assertEqual(result['state']['record_total'],103)
        self.assertEqual(result['state']['record_offset'],100)
        self.assertEqual(result['rows'][-1]['cells'][6]['value'],'Last record correction')
        self.assertNotIn('record-0:6',result['state']['field_decisions'])
        first=self.request('/api/workspace?job='+state['id'])
        self.assertEqual(len(first['rows']),50)
        self.assertEqual(first['rows'][0]['cells'][6]['value'],'First record correction')
        payload=app.export_data(state)
        self.assertEqual(len(payload['rows']),103)
        self.assertEqual(payload['rows'][-1][6],'Last record correction')

    def test_large_export_preserves_corrections_formats_markers_and_literal_text(self):
        from batch_export import write_workbook
        state=fixture()
        state.update(schema_version=2,created_date='10/11/2026')
        state['field_decisions']={'record-1:18':'903210','record-1:19':'474','record-1:20':'6.87',
                                  'record-1:8':'00123','record-1:6':'=2+2','record-1:4':'col.'}
        target=app.DATA/'export.xlsx'
        write_workbook(target,state,app.export_data,app.HEADERS)
        book=load_workbook(target)
        sheet=book['Converted records']
        self.assertEqual(sheet.cell(2,19).value,903210)
        self.assertIn('$',sheet.cell(2,19).number_format)
        self.assertEqual(sheet.cell(2,20).value,474)
        self.assertAlmostEqual(sheet.cell(2,21).value,.0687)
        self.assertEqual(sheet.cell(2,9).value,'00123')
        self.assertEqual(sheet.cell(2,7).value,'=2+2')
        self.assertNotEqual(sheet.cell(2,7).data_type,'f')
        self.assertEqual(sheet.cell(2,5).fill.fgColor.rgb[-6:],'E2F2E8')
        self.assertEqual(sheet.cell(2,11).fill.fgColor.rgb[-6:],'FDE2DE')
        self.assertTrue(any('[0/O]' in str(row) for row in book['Text regions'].values))
        self.assertTrue(any('Confirmed value; rule warning retained' in str(row) for row in book['Review status'].values))
        self.assertEqual(sheet.freeze_panes,'C2')
        book.close()


if __name__=='__main__':unittest.main()
