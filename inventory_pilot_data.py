"""Read-only inventory of source data; references are never used by scanning."""
from pathlib import Path
import hashlib,json,collections
from PIL import Image
import openpyxl

root=Path('D:/projects/piXeL');out=Path(__file__).resolve().parent/'outputs/pilot-testing';out.mkdir(parents=True,exist_ok=True)
images=[];workbooks=[]
for p in root.rglob('*'):
    if not p.is_file():continue
    if p.suffix.lower() in ('.jpg','.png','.jpeg'):
        item=dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest())
        try:
            with Image.open(p) as im: im.load();item.update(readable=True,size=list(im.size),format=im.format)
        except Exception as e:item.update(readable=False,error=str(e))
        images.append(item)
    elif p.suffix.lower()=='.xlsx':
        item=dict(path=str(p),sheets=[])
        try:
            w=openpyxl.load_workbook(p,read_only=True,data_only=True)
            for s in w:
                ids=[];examples=[]
                for row in s.values:
                    if len(examples)<2:examples.append(list(row)[:8])
                    for value in row[:3]:
                        if str(value).isdigit() and len(str(value))==10:ids.append(str(value));break
                item['sheets'].append(dict(name=s.title,rows=s.max_row,columns=s.max_column,record_ids=sorted(set(ids)),first_rows=examples))
        except Exception as e:item['error']=str(e)
        workbooks.append(item)
payload=dict(images=images,workbooks=workbooks)
(out/'inventory.json').write_text(json.dumps(payload,indent=2,default=str),encoding='utf-8')
print(json.dumps(dict(image_groups=[dict(folder=k,total=len(v),readable=sum(x['readable'] for x in v)) for k,v in ((k,[i for i in images if str(Path(i['path']).parent)==k]) for k in sorted({str(Path(i['path']).parent) for i in images}))],workbooks=[dict(path=w['path'],sheets=[{k:v for k,v in s.items() if k!='record_ids'}|dict(record_ids=len(s['record_ids']),range=[s['record_ids'][0],s['record_ids'][-1]] if s['record_ids'] else []) for s in w['sheets']]) for w in workbooks]),default=str,indent=2))
