import copy,io,json,base64,tempfile,threading,unittest,urllib.request,urllib.error
from pathlib import Path
from unittest.mock import patch
from PIL import Image,ImageDraw
import review_app as app
from pilot_rules import apply_rule,adjacent_pages,HEADERS
from detect_records import blocks
from test_sheet_review import fixture
from openpyxl import load_workbook

class RuleTests(unittest.TestCase):
    def test_numeric_choices_and_leading_zeros(self):
        value,issues=apply_rule(8,'00[0/O]12')
        self.assertEqual(value,'00012');self.assertTrue(issues)
        self.assertEqual(apply_rule(8,'00O12')[0],'00O12')
        self.assertTrue(apply_rule(8,'00O12')[1])
        self.assertEqual(apply_rule(8,'[1/7]001')[0],'[1/7]001')
    def test_safe_formatting_and_no_rounding(self):
        self.assertEqual(apply_rule(18,'903210')[0],'$903,210.00')
        self.assertEqual(apply_rule(19,'474')[0],'474 Months')
        self.assertEqual(apply_rule(20,'6.87')[0],'6.87%')
        self.assertEqual(apply_rule(18,'1.234')[0],'$1.23')
        self.assertFalse(apply_rule(18,'1.234')[1])
        self.assertEqual(apply_rule(34,'Amnerican Express')[0],'Amnerican Express Credit Card')
        self.assertEqual(apply_rule(34,'Amnerican Express Credit Card')[0],'Amnerican Express Credit Card')
    def test_date_and_property_are_not_guessed(self):
        self.assertEqual(apply_rule(9,'Tuesday, September 16 1952')[0],'Tuesday, 09/16/1952')
        self.assertTrue(apply_rule(9,'Turcday, Satmber 16 1952')[1])
        self.assertTrue(apply_rule(9,'02/30/1952')[1])
        self.assertEqual(apply_rule(15,'95279148317-2465')[0],'95279148317-2465')
    def test_page_gaps_and_reversed_order(self):
        self.assertTrue(adjacent_pages('AQ159_enc.jpg','AQ160_enc.jpg'))
        self.assertFalse(adjacent_pages('AQ159_enc.jpg','AQ161_enc.jpg'))
        self.assertFalse(adjacent_pages('AQ160_enc.jpg','AQ159_enc.jpg'))
    def test_old_schema_stays_unchanged(self):
        state=fixture();before=copy.deepcopy(state);sheet=app.sheet_data(state)
        self.assertIsNone(sheet[0]['cells'][2]['value']);self.assertEqual(state,before)
    def test_confirmed_name_conflict_stays_green_with_warning_and_remarks(self):
        state=fixture();state.update(schema_version=2,created_date='09/29/2026')
        state['field_decisions']={'record-1:4':'Mr.','record-1:5':'Clilnton Diblrmagli.','record-1:8':'00123'}
        row=app.sheet_data(state)[0]['cells']
        self.assertEqual(row[4]['status'],'confirmed');self.assertEqual(row[5]['status'],'confirmed')
        self.assertFalse(row[4]['rule_issues']);self.assertTrue(row[5]['rule_issues'])
        self.assertEqual(row[13]['value'],'ALN IS INVALID.')
        self.assertEqual(row[8]['value'],'00123');self.assertEqual(row[2]['value'],'09/29/2026')
        self.assertEqual(state['field_decisions']['record-1:4'],'Mr.')
        payload=app.export_data(state)
        self.assertEqual(payload['statuses'][0][4],'confirmed')
        warning=next(r for r in payload['review'] if r[1]==HEADERS[5] and r[2]=='Field rule')
        self.assertEqual(warning[-1],'Confirmed value; rule warning retained')
        del state['field_decisions']['record-1:4']
        self.assertNotEqual(app.sheet_data(state)[0]['cells'][4]['status'],'confirmed')

    def test_unconfirmed_rule_warning_still_requires_review(self):
        from pilot_rules import decorate
        state=fixture();state.update(schema_version=2,created_date='10/01/2026')
        cells=[dict(field=i,value='[?]',plain='[?]',status='unmapped',manual=False) for i in range(37)]
        cells[4].update(value='col.',plain='col.',status='read')
        result=decorate(state,state['records'][0],cells)
        self.assertEqual(result[4]['status'],'uncertain')
        self.assertTrue(result[4]['rule_issues'])
        self.assertEqual(result[4]['value'],'Col.')
    def test_single_column_final_page(self):
        p=Path('D:/projects/piXeL/FORM TTYFRD_8AQ/FORM TTYFRD_8AQ398_enc.jpg')
        if not p.exists():self.skipTest('Source absent')
        self.assertEqual(len(blocks(Image.open(p))),4)
    def test_boundary_edits_preserve_every_fragment_and_character(self):
        state=fixture();record=state['records'][0]
        record['fragments']=[dict(id='a',status='record_start'),dict(id='b',status='continuation_inferred',header_ocr='unknown')]
        record['lines'][0]['fragment']='a';record['lines'][1]['fragment']='b'
        state['decisions']={'l2s0':'L'}
        original_lines=copy.deepcopy(record['lines'])
        app.change_boundary(state,'record-1','split_here','b')
        self.assertEqual(len(state['records']),2)
        app.change_boundary(state,state['records'][1]['key'],'join_previous')
        self.assertEqual(len(state['records']),1)
        self.assertEqual(state['records'][0]['lines'],original_lines)
        self.assertEqual(state['decisions'],{'l2s0':'L'})
        state['field_decisions']={'record-1:6':'Confirmed address'}
        with self.assertRaises(ValueError):app.change_boundary(state,'record-1','split_here','b')
    def test_short_text_fragment_is_retained(self):
        im=Image.new('RGB',(600,300),'white');d=ImageDraw.Draw(im)
        for left,right in ((10,290),(310,590)):
            for x in (left,right):d.line((x,10,x,280),fill='black',width=2)
            for y in (10,40,280):d.line((left,y,right,y),fill='black',width=2)
            d.rectangle((left+20,22,left+40,29),fill='black')
        found=blocks(im);self.assertEqual(sum(b['short_fragment'] for b in found),2)

class ApiTests(unittest.TestCase):
    def test_scan_continuation_and_missing_page_keep_fragments(self):
        original=app.DATA
        with tempfile.TemporaryDirectory(prefix='pixel-fragment-test-') as root:
            app.DATA=Path(root)
            try:
                for number,second,expected in ((1,'AQ2.jpg',1),(2,'AQ3.jpg',2)):
                    state=dict(id=str(number)*32,status='scanning',records=[],decisions={},history=[],pages=[dict(name='AQ1.jpg',file='page-0.png'),dict(name=second,file='page-1.png')])
                    directory=app.folder(state['id']);directory.mkdir()
                    for page in state['pages']:Image.new('RGB',(100,100),'white').save(directory/page['file'])
                    block=dict(box=[0,0,100,100],column=0,block=0,first_in_column=True,open_bottom=True)
                    with patch('rapidocr_onnxruntime.RapidOCR',return_value=lambda *a,**kw:([],0)),patch('detect_records.blocks',return_value=[block]),patch.object(app,'read_header',side_effect=[('1234567890','1234567890'),(None,'continuation')]):
                        app.scan(state)
                    self.assertEqual(state['status'],'ready');self.assertEqual(len(state['records']),expected)
                    self.assertEqual(sum(len(r['fragments']) for r in state['records']),2)
                    self.assertTrue(all('No text detected in this block.' in r['notes'] for r in state['records']))
                    if expected==2:self.assertTrue(state['page_issues'])
            finally:app.DATA=original

    def test_project_validation_and_highlighted_export(self):
        original=app.DATA
        with tempfile.TemporaryDirectory(prefix='pixel-pilot-test-') as root:
            app.DATA=Path(root);server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
            t=threading.Thread(target=server.serve_forever,daemon=True);t.start()
            def request(path,data=None):
                req=urllib.request.Request(f'http://127.0.0.1:{server.server_port}'+path,data=None if data is None else json.dumps(data).encode(),headers={'Content-Type':'application/json','X-Review-Token':app.TOKEN})
                with urllib.request.urlopen(req,timeout=100) as res:return res.read()
            try:
                raw=io.BytesIO();Image.new('RGB',(100,100),'white').save(raw,format='PNG')
                page={'name':'AQ1.jpg','data':base64.b64encode(raw.getvalue()).decode()}
                with self.assertRaises(urllib.error.HTTPError):request('/api/scan',{'name':'Duplicate','pages':[page,page]})
                self.assertFalse(app.list_jobs())
                self.assertFalse([p for p in app.DATA.iterdir() if p.name!='_diagnostics'])
                with patch.object(app,'scan') as scan:
                    result=json.loads(request('/api/scan',{'name':'My pilot','pages':[page]}))
                    state=app.load(result['id']);self.assertEqual(state['project_name'],'My pilot');self.assertEqual(state['schema_version'],2)
                state=fixture();state.update(schema_version=2,project_name='Export fixture',created_date='09/29/2026')
                directory=app.folder(state['id']);directory.mkdir()
                state['records'][0]['fragments']=[dict(id='p0-b0',page='page.jpg',file='page-0.png',box=[0,0,100,100],column=0,status='record_start')]
                state['field_decisions']={'record-1:18':'903210','record-1:19':'474','record-1:20':'6.87','record-1:8':'00123','record-1:6':'=2+2','record-1:4':'col.'}
                app.save(state)
                result=json.loads(request('/api/export',{'job':state['id']}));book=load_workbook(io.BytesIO(request(result['url'])))
                sheet=book['Converted records'];self.assertEqual(sheet['A1'].value,HEADERS[0]);self.assertEqual(sheet.cell(2,19).value,903210)
                self.assertIn('$',sheet.cell(2,19).number_format);self.assertEqual(sheet.cell(2,20).value,474)
                self.assertAlmostEqual(sheet.cell(2,21).value,.0687);self.assertEqual(sheet.cell(2,9).value,'00123')
                self.assertNotEqual(sheet.cell(2,7).data_type,'f');self.assertEqual(sheet.cell(2,2).value,'1234567890')
                self.assertTrue(any('[0/O]' in str(r) for r in book['Text regions'].values))
                self.assertTrue(any('Field-rule alternative filter' in str(r) for r in book['Review status'].values))
                self.assertEqual(sheet.cell(2,19).fill.fgColor.rgb[-6:],'E2F2E8')
                self.assertEqual(sheet.cell(2,5).value,'Col.')
                self.assertEqual(sheet.cell(2,5).fill.fgColor.rgb[-6:],'E2F2E8')
                rule_rows=[r for r in book['Review status'].values if r[1]==HEADERS[4] and r[2]=='Field rule']
                self.assertEqual(rule_rows[0][-1],'Confirmed value; rule warning retained')
                self.assertEqual(sheet.cell(2,6).fill.fgColor.rgb[-6:],'FDE2DE')
                self.assertEqual(sheet.cell(2,11).fill.fgColor.rgb[-6:],'FDE2DE')
                self.assertEqual(book['Source fragments'].max_row,2);book.close()
            finally:server.shutdown();server.server_close();t.join();app.DATA=original

if __name__=='__main__':unittest.main()
